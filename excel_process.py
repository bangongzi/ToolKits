import openpyxl

##### 读 #####
workbook = openpyxl.load_workbook('测试表.xlsx')
sheet = workbook['sheet1']
print(sheet.max_column, sheet.max_row)  # 这里我们可以看到这个sheet有多少行，多少列

# 假定第一列有值，把第一列的值复制到第二列去
row_index = 1
while row_index <= sheet.max_row:
    value = sheet.cell(row_index, 1).value
    sheet.cell(row_index, 2, value)
    row_index = row_index + 1

workbook.save('测试表.xlsx')

##### 写 #####
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = 'sheet1'
sheet.cell(row=1, column=1, value='上金刚')  # 第一行第一列
sheet.cell(row=2, column=1, value='懒扎衣')  # 第二行第一列
workbook.save('测试表.xlsx')
